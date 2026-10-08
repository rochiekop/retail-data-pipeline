import pytest
from openpyxl import Workbook, load_workbook

from pipeline.bronze import extract_catalog
from pipeline.olist import CATALOG_SHEETS, SourceSchemaError


def _rows(client, query):
    return client.query(query).result_rows


def test_loads_every_sheet_with_its_checksum(catalog_file, warehouse_client):
    assert extract_catalog(catalog_file) == 5  # 3 products + 2 translations
    assert _rows(
        warehouse_client,
        "select product_category_name, length(_source_checksum) from bronze.products where product_id = 'p1'",
    ) == [("brinquedos", 64)]


def test_unchanged_catalog_is_not_reloaded(catalog_file, warehouse_client):
    extract_catalog(catalog_file)
    assert extract_catalog(catalog_file) is None
    assert _rows(warehouse_client, "select count() from bronze.products") == [(3,)]


def test_edited_catalog_is_loaded_again(catalog_file, warehouse_client):
    extract_catalog(catalog_file)
    workbook = load_workbook(catalog_file)
    workbook["products"]["B3"] = "brinquedos"  # p2 gets a category
    workbook.save(catalog_file)
    assert extract_catalog(catalog_file) == 5
    assert _rows(warehouse_client, "select uniqExact(_loaded_at) from bronze.products") == [(2,)]


def test_emptied_sheet_replaces_the_older_load(catalog_file, warehouse_client):
    extract_catalog(catalog_file)
    workbook = load_workbook(catalog_file)
    workbook["category_translation"].delete_rows(2, 2)  # every translation removed
    workbook.save(catalog_file)
    assert extract_catalog(catalog_file) == 3
    latest_load = "(select max(_loaded_at) from bronze.catalog_loads)"
    assert _rows(
        warehouse_client,
        f"select (select count() from bronze.products where _loaded_at = {latest_load}), "
        f"(select count() from bronze.category_translation where _loaded_at = {latest_load})",
    ) == [(3, 0)]


def test_blank_rows_are_ignored(catalog_file, warehouse_client):
    workbook = load_workbook(catalog_file)
    workbook["products"].append([None] * 9)
    workbook.save(catalog_file)
    assert extract_catalog(catalog_file) == 5


def test_trailing_blank_cells_load_as_null(tmp_path, warehouse_client):
    # A write-only workbook (as build_catalog makes) stores no cells after a row's last value,
    # so this row reads back as ("p9",). Olist has two products like it.
    path = tmp_path / "catalog.xlsx"
    workbook = Workbook(write_only=True)
    for sheet, (_, columns) in CATALOG_SHEETS.items():
        workbook.create_sheet(sheet).append(columns)
    workbook.worksheets[0].append(["p9"])
    workbook.save(path)
    assert extract_catalog(path) == 1
    assert _rows(
        warehouse_client, "select product_category_name, product_width_cm from bronze.products where product_id = 'p9'"
    ) == [(None, None)]


def test_numbers_typed_into_excel_are_stored_as_text(catalog_file, warehouse_client):
    workbook = load_workbook(catalog_file)
    workbook["products"]["F2"] = 225  # a number, as a person editing Excel would type it
    workbook.save(catalog_file)
    extract_catalog(catalog_file)
    assert _rows(
        warehouse_client, "select product_weight_g from bronze.products where product_id = 'p1'"
    ) == [("225",)]


def test_missing_catalog_fails_the_run(tmp_path, warehouse_client):
    with pytest.raises(FileNotFoundError, match="Merchandising Catalog"):
        extract_catalog(tmp_path / "missing.xlsx")


def test_changed_columns_fail_the_run(catalog_file, warehouse_client):
    workbook = load_workbook(catalog_file)
    workbook["products"]["B1"] = "category"
    workbook.save(catalog_file)
    with pytest.raises(SourceSchemaError, match=r"merchandising_catalog.xlsx\[products\]"):
        extract_catalog(catalog_file)
    catalog_file.unlink()  # the workbook must be closed after a failed check


def test_missing_sheet_fails_the_run(tmp_path, warehouse_client):
    path = tmp_path / "catalog.xlsx"
    workbook = Workbook()
    workbook.active.title = "products"
    workbook.save(path)
    with pytest.raises(SourceSchemaError, match="category_translation"):
        extract_catalog(path)
