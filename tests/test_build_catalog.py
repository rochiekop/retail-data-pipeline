import shutil

import pytest
from openpyxl import load_workbook

from pipeline.olist import SourceSchemaError
from pipeline.seed import build_catalog
from tests.conftest import FIXTURES


def _rows(path, sheet):
    return [list(r) for r in load_workbook(path, read_only=True)[sheet].iter_rows(values_only=True)]


def test_catalog_has_one_sheet_per_catalog_table(catalog_file):
    assert load_workbook(catalog_file).sheetnames == ["products", "category_translation"]


def test_header_ignores_the_byte_order_mark(catalog_file):
    assert _rows(catalog_file, "category_translation")[0] == [
        "product_category_name",
        "product_category_name_english",
    ]


def test_values_are_copied_unmodified(catalog_file):
    products = _rows(catalog_file, "products")
    assert products[1] == ["p1", "brinquedos", "40", "287", "1", "225", "16", "10", "14"]
    assert products[2][:2] == ["p2", None]


def test_output_directory_is_created(tmp_path):
    out = tmp_path / "nested" / "catalog.xlsx"
    build_catalog(FIXTURES, out)
    assert out.exists()


def test_changed_columns_fail_the_build(tmp_path):
    source = tmp_path / "olist"
    shutil.copytree(FIXTURES, source)
    products = source / "olist_products_dataset.csv"
    header_renamed = products.read_text(encoding="utf-8").replace("product_category_name", "category", 1)
    products.write_text(header_renamed, encoding="utf-8")
    with pytest.raises(SourceSchemaError, match="olist_products_dataset.csv"):
        build_catalog(source, tmp_path / "catalog.xlsx")
