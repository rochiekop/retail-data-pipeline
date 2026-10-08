"""Bronze: extract each source and load an unmodified, text-typed copy into the warehouse."""

import hashlib
from datetime import date, datetime, timezone
from io import BytesIO
from pathlib import Path

import psycopg
from clickhouse_connect.driver.client import Client
from openpyxl import load_workbook

from pipeline import config
from pipeline.olist import CATALOG_SHEETS, SHOP_TABLES, SourceSchemaError
from pipeline.warehouse import connect


def _select_list(table: str, alias: str = "x") -> str:
    return ", ".join(f"{alias}.{c}::text" for c in SHOP_TABLES[table][1])


# Business Date: the date part of the purchase timestamp as stored (Sao Paulo local time).
_ON_BUSINESS_DATE = "o.order_purchase_timestamp::date = %(d)s"

SHOP_EXTRACTS: dict[str, str] = {
    "orders": f"select {_select_list('orders', 'o')} from shop.orders o where {_ON_BUSINESS_DATE}",
    "order_items": f"select {_select_list('order_items')} from shop.order_items x "
    f"join shop.orders o on o.order_id = x.order_id where {_ON_BUSINESS_DATE}",
    "order_payments": f"select {_select_list('order_payments')} from shop.order_payments x "
    f"join shop.orders o on o.order_id = x.order_id where {_ON_BUSINESS_DATE}",
}

_SHOP_TABLE_DDL = """
create table if not exists bronze.{table} (
    {columns},
    _business_date Date,
    _loaded_at DateTime64(6, 'UTC')
)
engine = MergeTree
partition by toYYYYMM(_business_date)
order by (_business_date, _loaded_at)
"""

# One row per shop extract, written after its tables, even when it found no rows. Silver reads every
# shop table at the latest of these, so a rerun that finds nothing replaces the older load.
_SHOP_LOADS_DDL = """
create table if not exists bronze.shop_loads (
    _business_date Date,
    _loaded_at DateTime64(6, 'UTC'),
    row_counts Map(String, UInt64)
)
engine = MergeTree
order by (_business_date, _loaded_at)
"""

_CATALOG_TABLE_DDL = """
create table if not exists bronze.{table} (
    {columns},
    _source_checksum String,
    _loaded_at DateTime64(6, 'UTC')
)
engine = MergeTree
order by _loaded_at
"""

# One row per loaded catalog file, written after both sheets, even when a sheet is empty. It is what
# "already loaded" checks, and Silver reads both sheets at the latest of these.
_CATALOG_LOADS_DDL = """
create table if not exists bronze.catalog_loads (
    _source_checksum String,
    _loaded_at DateTime64(6, 'UTC'),
    row_counts Map(String, UInt64)
)
engine = MergeTree
order by _loaded_at
"""


def _text_columns(columns: list[str]) -> str:
    return ",\n    ".join(f"{c} Nullable(String)" for c in columns)


def ensure_bronze_tables(client: Client) -> None:
    client.command("create database if not exists bronze")
    for table in SHOP_EXTRACTS:
        _, columns = SHOP_TABLES[table]
        client.command(_SHOP_TABLE_DDL.format(table=table, columns=_text_columns(columns)))
    client.command(_SHOP_LOADS_DDL)
    for sheet, (_, columns) in CATALOG_SHEETS.items():
        client.command(_CATALOG_TABLE_DDL.format(table=sheet, columns=_text_columns(columns)))
    client.command(_CATALOG_LOADS_DDL)


def extract_shop(business_date: date) -> dict[str, int]:
    # ClickHouse has no multi-table transactions: each table's insert is atomic on its own, and the
    # load only becomes visible to Silver once bronze.shop_loads records it, after every table.
    loaded_at = datetime.now(timezone.utc)
    client = connect()
    try:
        ensure_bronze_tables(client)
        counts = {}
        with psycopg.connect(config.shop_db().conninfo()) as shop:
            for table, query in SHOP_EXTRACTS.items():
                rows = shop.execute(query, {"d": business_date}).fetchall()
                if rows:
                    client.insert(
                        table,
                        [(*row, business_date, loaded_at) for row in rows],
                        column_names=[*SHOP_TABLES[table][1], "_business_date", "_loaded_at"],
                        database="bronze",
                    )
                counts[table] = len(rows)
        client.insert(
            "shop_loads",
            [(business_date, loaded_at, counts)],
            column_names=["_business_date", "_loaded_at", "row_counts"],
            database="bronze",
        )
        return counts
    finally:
        client.close()


def _read_catalog(name: str, content: bytes) -> dict[str, list[tuple]]:
    workbook = load_workbook(BytesIO(content), read_only=True)
    try:
        missing = [sheet for sheet in CATALOG_SHEETS if sheet not in workbook.sheetnames]
        if missing:
            raise SourceSchemaError(f"{name}: missing sheets {missing}")
        sheets = {}
        for sheet, (_, columns) in CATALOG_SHEETS.items():
            rows = workbook[sheet].iter_rows(values_only=True)
            header = list(next(rows, ()))
            if header != columns:
                raise SourceSchemaError(f"{name}[{sheet}]: expected columns {columns}, got {header}")
            # A workbook can store no cells after a row's last value (build_catalog's do), so pad to the header width.
            sheets[sheet] = [
                tuple(None if v is None else str(v) for v in (row + (None,) * len(columns))[: len(columns)])
                for row in rows
                if any(v is not None for v in row)
            ]
        return sheets
    finally:
        workbook.close()


def extract_catalog(path: Path) -> int | None:
    if not path.exists():
        raise FileNotFoundError(f"Merchandising Catalog not found: {path}")
    # Read the file once, so the checksum always describes the rows that are loaded.
    content = path.read_bytes()
    checksum = hashlib.sha256(content).hexdigest()
    sheets = _read_catalog(path.name, content)

    loaded_at = datetime.now(timezone.utc)
    client = connect()
    try:
        ensure_bronze_tables(client)
        already_loaded = client.query(
            "select 1 from bronze.catalog_loads where _source_checksum = {c:String} limit 1",
            parameters={"c": checksum},
        ).result_rows
        if already_loaded:
            return None
        # Two parallel runs may both load a new file; Silver reads only the latest load, so that is harmless.
        for sheet, rows in sheets.items():
            if rows:
                client.insert(
                    sheet,
                    [(*row, checksum, loaded_at) for row in rows],
                    column_names=[*CATALOG_SHEETS[sheet][1], "_source_checksum", "_loaded_at"],
                    database="bronze",
                )
        row_counts = {sheet: len(rows) for sheet, rows in sheets.items()}
        client.insert(
            "catalog_loads",
            [(checksum, loaded_at, row_counts)],
            column_names=["_source_checksum", "_loaded_at", "row_counts"],
            database="bronze",
        )
        return sum(row_counts.values())
    finally:
        client.close()
